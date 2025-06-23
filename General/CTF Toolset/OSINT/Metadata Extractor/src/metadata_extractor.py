import os
import sys
import argparse

from PIL import Image
import exifread
import PyPDF2
import docx
import openpyxl
import pptx

def extract_image_metadata(file_path):
    # Extract EXIF metadata from image files
    metadata = {}
    with open(file_path, 'rb') as f:
        tags = exifread.process_file(f, details=False)
        for tag in tags:
            metadata[tag] = str(tags[tag])
    # Extract GPS data if present
    gps_info = {k: v for k, v in metadata.items() if "GPS" in k}
    if gps_info:
        metadata['GPS_Info'] = gps_info
    return metadata

def extract_pdf_metadata(file_path):
    # Extract metadata from PDF files
    metadata = {}
    with open(file_path, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        if reader.metadata:
            for k, v in reader.metadata.items():
                metadata[k] = str(v)
    return metadata

def extract_docx_metadata(file_path):
    # Extract core properties from DOCX files
    metadata = {}
    doc = docx.Document(file_path)
    props = doc.core_properties
    metadata.update({
        "author": props.author,
        "title": props.title,
        "created": str(props.created),
        "last_modified_by": props.last_modified_by,
        "modified": str(props.modified)
    })
    # Remove empty values
    return {k: v for k, v in metadata.items() if v}

def extract_xlsx_metadata(file_path):
    # Extract metadata from XLSX Excel files
    metadata = {}
    wb = openpyxl.load_workbook(file_path)
    props = wb.properties
    metadata.update({
        "author": props.creator,
        "title": props.title,
        "created": str(props.created),
        "modified": str(props.modified),
        "last_modified_by": props.lastModifiedBy
    })
    return {k: v for k, v in metadata.items() if v}

def extract_pptx_metadata(file_path):
    # Extract metadata from PPTX presentation files
    metadata = {}
    pres = pptx.Presentation(file_path)
    props = pres.core_properties
    metadata.update({
        "author": props.author,
        "title": props.title,
        "created": str(props.created),
        "modified": str(props.modified),
        "last_modified_by": props.last_modified_by
    })
    return {k: v for k, v in metadata.items() if v}

def extract_metadata(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    # Route extraction to appropriate handler based on file extension
    if extension in ['.jpg', '.jpeg', '.png']:
        return extract_image_metadata(file_path)
    elif extension == '.pdf':
        return extract_pdf_metadata(file_path)
    elif extension == '.docx':
        return extract_docx_metadata(file_path)
    elif extension == '.xlsx':
        return extract_xlsx_metadata(file_path)
    elif extension == '.pptx':
        return extract_pptx_metadata(file_path)
    else:
        return {"error": "Unsupported file type"}

def print_metadata(file_path, metadata):
    print(f"\n--- Metadata for {file_path} ---")
    if "error" in metadata:
        print("Error:", metadata["error"])
        return
    for key, value in metadata.items():
        if isinstance(value, dict):
            print(f"{key}:")
            for sub_k, sub_v in value.items():
                print(f"  {sub_k}: {sub_v}")
        else:
            print(f"{key}: {value}")

def main():
    parser = argparse.ArgumentParser(description="Extract metadata from files")
    parser.add_argument("--file", help="Path to file")
    parser.add_argument("--dir", help="Directory of files")

    args = parser.parse_args()

    files_to_scan = []

    if args.file:
        files_to_scan.append(args.file)
    elif args.dir:
        # Scan all files in given directory
        for filename in os.listdir(args.dir):
            full_path = os.path.join(args.dir, filename)
            if os.path.isfile(full_path):
                files_to_scan.append(full_path)
    else:
        print("Specify either --file or --dir")
        sys.exit(1)

    for file_path in files_to_scan:
        metadata = extract_metadata(file_path)
        print_metadata(file_path, metadata)

if __name__ == "__main__":
    main()
